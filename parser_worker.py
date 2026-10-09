"""Short-lived parser: stdin bytes, bounded JSON stdout, no upload files or network."""
import base64
import io
import json
import sys
import warnings
import os
import zipfile
import math
import re
from itertools import zip_longest
from datetime import date, datetime, time
from pathlib import PurePosixPath

MAX_STREAM = 2 * 1024 * 1024

def check_xlsx(raw):
    from defusedxml.ElementTree import iterparse
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            parts=archive.infolist()
            if not 1<=len(parts)<=256: raise ValueError('Excel 内部文件数量超过限制')
            names=set(); total=0
            for part in parts:
                path=PurePosixPath(part.filename)
                if part.filename in names or path.is_absolute() or '..' in path.parts or '\\' in part.filename:
                    raise ValueError('Excel 压缩包结构无效')
                names.add(part.filename)
                lower=part.filename.lower()
                if 'vbaproject' in lower or lower.startswith(('xl/embeddings/','xl/externallinks/')):
                    raise ValueError('请移除宏、嵌入对象或外部工作簿链接后保存为 .xlsx')
                if part.flag_bits & 1 or part.compress_type not in (zipfile.ZIP_STORED,zipfile.ZIP_DEFLATED):
                    raise ValueError('Excel 压缩或加密格式不支持')
                total+=part.file_size
                limit=4*1024*1024 if lower=='xl/sharedstrings.xml' else 8*1024*1024
                if part.file_size>limit or total>24*1024*1024:
                    raise ValueError('Excel 解压后内容过大，请精简工作簿')
                # Read without extracting. CRC and actual output length are checked before parsing.
                with archive.open(part) as source: content=source.read(limit+1)
                if len(content)!=part.file_size: raise ValueError('Excel 内部文件长度无效')
                if lower.endswith(('.xml','.rels')):
                    normalized=content.replace(b'\x00',b'').upper()
                    if b'<!DOCTYPE' in normalized or b'<!ENTITY' in normalized:
                        raise ValueError('Excel 不允许 XML 实体或外部文档声明')
                if lower=='[content_types].xml' and b'macroenabled' in content.lower():
                    raise ValueError('只支持无宏的 .xlsx 文件')
                if lower.startswith('xl/worksheets/') and lower.endswith('.xml'):
                    previous_row=0; current_row=0; previous_column=0
                    for event,node in iterparse(io.BytesIO(content),events=('start','end'),forbid_dtd=True):
                        tag=node.tag.rsplit('}',1)[-1]
                        if event=='start' and tag=='row':
                            current_row=int(node.get('r','0'))
                            if not previous_row<current_row<=10001: raise ValueError('工作表行号无效或超过 10,001 行（含表头）')
                            previous_row=current_row; previous_column=0
                        if event=='start' and tag=='c':
                            coordinate=re.fullmatch(r'([A-Z]{1,3})([1-9][0-9]*)',node.get('r',''))
                            if not coordinate: raise ValueError('Excel 单元格坐标无效')
                            column=0
                            for letter in coordinate[1]: column=column*26+ord(letter)-64
                            if not previous_column<column<=64 or int(coordinate[2])!=current_row:
                                raise ValueError('工作表最多 64 列，且单元格坐标须按行列排序')
                            previous_column=column
                        if event=='end': node.clear()
            if not {'[Content_Types].xml','xl/workbook.xml'}.issubset(names):
                raise ValueError('不是有效的 .xlsx 工作簿')
    except (zipfile.BadZipFile,RuntimeError,NotImplementedError):
        raise ValueError('Excel 文件损坏或格式不支持') from None

def xlsx_data(raw,options):
    check_xlsx(raw)
    os.environ['OPENPYXL_LXML']='False'
    os.environ['OPENPYXL_DEFUSEDXML']='True'
    from openpyxl import load_workbook
    options=options or {}; selected=options.get('sheet')
    columns=options.get('columns',[])
    book=load_workbook(io.BytesIO(raw),read_only=True,data_only=False,keep_links=False)
    cached=None
    try:
        sheets=[ws for ws in book.worksheets if ws.sheet_state=='visible']
        if not 1<=len(book.worksheets)<=20 or not sheets: raise ValueError('Excel 限 1–20 张工作表且需有可见数据表')
        names=[ws.title for ws in sheets]
        if selected is None: return {'sheets':names}
        if type(selected) is not int or not 0<=selected<len(sheets): raise ValueError('工作表选择无效')
        if not columns or len(columns)>64: raise ValueError('实验模板列无效')
        cached=load_workbook(io.BytesIO(raw),read_only=True,data_only=True,keep_links=False)
        worksheet=sheets[selected]; values=cached[worksheet.title]
        # Ignore producer-supplied dimensions; bound actual streamed rows and cells instead.
        worksheet.reset_dimensions(); values.reset_dimensions()
        headers=None; rows=[]; formula_count=0
        for number,pair in enumerate(zip_longest(worksheet.iter_rows(),values.iter_rows()),1):
            actual,saved=pair
            if number>10001 or actual is None or saved is None: raise ValueError('工作表最多 10,001 行（含表头），请精简后导入')
            if max(len(actual),len(saved))>64: raise ValueError('工作表最多 64 列，请移除多余列或格式')
            if headers is None:
                if not any(c.value is not None for c in actual): continue
                if any(c.data_type=='f' for c in actual): raise ValueError('表头须为模板列名，不能使用公式')
                headers=[str(c.value or '').strip().lstrip('\ufeff') for c in actual]
                while headers and not headers[-1]: headers.pop()
                if any(not h or len(h)>256 for h in headers) or len(set(headers))!=len(headers):
                    raise ValueError('Excel 列名不可重复或为空，最多 256 字符')
                missing=[c for c in columns if c not in headers]
                if missing: raise ValueError('缺少列：'+', '.join(missing))
                continue
            if not any(c.value is not None for c in actual): continue
            if any(c.value is not None for c in actual[len(headers):]): raise ValueError('数据列超出表头范围')
            row={}
            for column in columns:
                index=headers.index(column)
                cell=actual[index] if index<len(actual) else None
                value=cell.value if cell else None
                if cell and cell.data_type=='f':
                    formula_count+=1
                    saved_cell=saved[index] if index<len(saved) else None
                    if saved_cell is None or saved_cell.value is None:
                        raise ValueError('公式缺少已保存的计算结果，请在 Excel 中重新计算并保存，或粘贴为数值后导入')
                    cell=saved_cell; value=cell.value
                if cell and cell.data_type=='e': raise ValueError('数据含 Excel 公式错误，请修正后导入')
                if isinstance(value,bool): raise ValueError('数据单元格不能是布尔值')
                if isinstance(value,float) and not math.isfinite(value): raise ValueError('数据含非有限数值')
                if isinstance(value,(date,datetime,time)): value=value.isoformat()
                value='' if value is None else str(value)
                if len(value)>256: raise ValueError('单元格最多 256 字符')
                row[column]=value
            rows.append(row)
            if len(rows)>10000: raise ValueError('Excel 最多 10,000 行数据')
        if not rows: raise ValueError('Excel 需要表头及至少一行数据')
        return {'sheets':names,'sheet':selected,'rows':rows,'formulaCount':formula_count}
    finally:
        book.close()
        if cached is not None: cached.close()

def pdf_text(raw):
    from pypdf import PdfReader, Configuration, apply_configuration
    config = Configuration(maximum_declared_stream_length=MAX_STREAM,
        array_based_stream_maximum_output_length=MAX_STREAM, zlib_maximum_output_length=MAX_STREAM,
        lzw_maximum_output_length=MAX_STREAM, run_length_maximum_output_length=MAX_STREAM,
        image_maximum_buffer_size=MAX_STREAM, page_tree_maximum_entries=100,
        page_tree_maximum_depth=20, xform_maximum_invocations_per_extraction=100,
        jbig2dec_binary=None)
    with apply_configuration(config):
        reader = PdfReader(io.BytesIO(raw), root_object_recovery_limit=1000)
        if reader.is_encrypted: raise ValueError('不接受加密 PDF，请在本地解密后导入')
        count = len(reader.pages)
        if not 1 <= count <= 20: raise ValueError('过程 PDF 限 1–20 页')
        chunks = []; total = 0; work = 0
        for page in reader.pages:
            content = page.get_contents()
            if content:
                work += len(content.get_data())
                if work > 4 * MAX_STREAM: raise ValueError('PDF 解压后内容过大，请改为粘贴文字')
            value = page.extract_text() or ''
            total += len(value)
            if total > 18000: raise ValueError('过程文字超过 18000 字符，请精简后导入')
            chunks.append(value)
        text = '\n'.join(chunks)
        if len(text.strip()) < 20: raise ValueError('PDF 文字不足或为扫描件，请手动填写')
        return {'text':text,'pages':count}

def image_data(raw):
    from PIL import Image
    Image.MAX_IMAGE_PIXELS = 2_000_000
    warnings.simplefilter('error', Image.DecompressionBombWarning)
    with Image.open(io.BytesIO(raw)) as image:
        if image.format not in ('PNG','JPEG') or image.width*image.height > 2_000_000 or max(image.size)>4096:
            raise ValueError('图片限 PNG/JPEG、200 万像素且边长不超过 4096')
        image.verify()
    with Image.open(io.BytesIO(raw)) as image:
        # Copy pixels into a new image: discard EXIF, comments and other source metadata.
        fresh = Image.new('RGB',image.size, 'white')
        if 'A' in image.getbands(): fresh.paste(image.convert('RGB'),mask=image.getchannel('A'))
        else: fresh.paste(image.convert('RGB'))
        target = io.BytesIO(); fresh.save(target,format='PNG')
        safe = target.getvalue()
        if len(safe)>4*1024*1024: raise ValueError('规范化图片超过 4 MB，请缩小图片')
        return {'data':'data:image/png;base64,'+base64.b64encode(safe).decode()}

def math_images(raw):
    # -I strips import paths; add only the trusted application directory.
    from pathlib import Path
    sys.path.insert(0,str(Path(__file__).resolve().parent))
    from math_layout import validate_tex
    from matplotlib.mathtext import MathTextParser
    from matplotlib.font_manager import FontProperties
    import numpy as np
    from PIL import Image
    values=json.loads(raw)
    if not isinstance(values,list) or not 1<=len(values)<=64: raise ValueError('每份文档最多 64 条公式')
    parser=MathTextParser('agg'); result=[]
    for tex in values:
        validate_tex(tex)
        try:
            parsed=parser.parse('$'+tex+'$',dpi=240,prop=FontProperties(size=12,math_fontfamily='stix'))
        except (ValueError,RuntimeError):
            raise ValueError('公式排版失败，请检查语法；PDF 支持分式、根号、上下标、积分等常用公式，不支持完整 LaTeX 环境') from None
        if parsed.width>1500 or parsed.height>400 or parsed.width*parsed.height>240000:
            raise ValueError('公式过宽或过高，请拆分为较短公式')
        alpha=np.asarray(parsed.image)
        image=Image.new('RGBA',(alpha.shape[1]+4,alpha.shape[0]+4),(0,0,0,0))
        ink=Image.new('RGBA',(alpha.shape[1],alpha.shape[0]),(24,37,29,255))
        ink.putalpha(Image.fromarray(alpha)); image.paste(ink,(2,2))
        target=io.BytesIO(); image.save(target,format='PNG')
        result.append({'data':base64.b64encode(target.getvalue()).decode(),
            'width':image.width*72/240,'height':image.height*72/240,'depth':(parsed.depth+2)*72/240})
    return {'formulas':result}

if __name__ == '__main__':
    try:
        raw=sys.stdin.buffer.read(8*1024*1024+1)
        if len(raw)>8*1024*1024: raise ValueError('文件过大')
        mode=sys.argv[1]
        if mode=='pdf': answer=pdf_text(raw)
        elif mode=='image': answer=image_data(raw)
        elif mode=='xlsx': answer=xlsx_data(raw,json.loads(sys.argv[2]))
        elif mode=='math': answer=math_images(raw)
        elif mode=='chart':
            from pathlib import Path
            sys.path.insert(0,str(Path(__file__).resolve().parent))
            from chart_import import import_chart
            answer=import_chart(raw,json.loads(sys.argv[2]))
        else: raise ValueError('文件解析模式无效')
        sys.stdout.write(json.dumps(answer,ensure_ascii=True))
    except ValueError as exc:
        sys.stdout.write(json.dumps({'error':str(exc)[:160]},ensure_ascii=True)); sys.exit(1)
    except Exception:
        sys.stdout.write(json.dumps({'error':'文件无效或超出解析限制，请检查格式或缩小文件'},ensure_ascii=True)); sys.exit(1)
