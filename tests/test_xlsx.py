import io
import sys
import unittest
import zipfile
from pathlib import Path
from openpyxl import Workbook

sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import server as s

COLUMNS=['relative_pressure','uptake_cm3_STP_g']

def workbook(rows=None):
    book=Workbook(); sheet=book.active; sheet.title='实验数据'
    for row in (rows or [COLUMNS,[.05,12.5],[.1,19]]): sheet.append(row)
    target=io.BytesIO();book.save(target);book.close();return target.getvalue()

def change_part(raw,name,transform):
    target=io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw)) as original, zipfile.ZipFile(target,'w',zipfile.ZIP_DEFLATED) as edited:
        for part in original.infolist():
            content=original.read(part.filename)
            edited.writestr(part.filename,transform(content) if part.filename==name else content)
    return target.getvalue()

def parse(raw,sheet=0):
    return s.isolated_parse('xlsx',raw,{'sheet':sheet,'columns':COLUMNS})

class XlsxTests(unittest.TestCase):
    def test_numeric_reordered_columns_blank_rows_and_extra_private_column(self):
        raw=workbook([['private_email',COLUMNS[1],COLUMNS[0]],['ignored@example.invalid',12.5,.05],[],['ignored',19,.1],['ignored',28,.2]])
        answer=parse(raw)
        self.assertEqual(answer['rows'],[{COLUMNS[0]:'0.05',COLUMNS[1]:'12.5'},{COLUMNS[0]:'0.1',COLUMNS[1]:'19'},{COLUMNS[0]:'0.2',COLUMNS[1]:'28'}])
        self.assertNotIn('ignored',str(answer)); self.assertEqual(answer['formulaCount'],0)
        data=s.analyze(3,answer['rows'])
        self.assertIn('results',data)

    def test_multiple_visible_sheets_and_selection(self):
        book=Workbook();book.active.title='说明';book.active.append(['使用说明'])
        sheet=book.create_sheet('BET 数据');sheet.append(COLUMNS);sheet.append([.1,19])
        hidden=book.create_sheet('隐藏资料');hidden.sheet_state='hidden';hidden.append(['不导入'])
        target=io.BytesIO();book.save(target);book.close();raw=target.getvalue()
        self.assertEqual(parse(raw,None),{'sheets':['说明','BET 数据']})
        self.assertEqual(parse(raw,1)['rows'][0][COLUMNS[1]],'19')
        with self.assertRaisesRegex(ValueError,'缺少列'):parse(raw,0)
        with self.assertRaises(ValueError):parse(raw,2)

    def test_formula_requires_cached_result_and_preserves_cached_value(self):
        raw=workbook([COLUMNS,[.1,'=1+2']])
        with self.assertRaisesRegex(ValueError,'计算结果'):parse(raw)
        raw=change_part(raw,'xl/worksheets/sheet1.xml',lambda content:content.replace(b'<f>1+2</f><v></v>',b'<f>1+2</f><v>3</v>').replace(b'<f>1+2</f><v />',b'<f>1+2</f><v>3</v>'))
        answer=parse(raw); self.assertEqual(answer['rows'][0][COLUMNS[1]],'3');self.assertEqual(answer['formulaCount'],1)

    def test_formula_error_missing_duplicate_headers_and_large_cells(self):
        for rows in [[COLUMNS,[.1,'#DIV/0!']],[['wrong',COLUMNS[1]],[.1,2]],[[COLUMNS[0],COLUMNS[0]],[.1,2]],[COLUMNS,[.1,'x'*257]]]:
            with self.assertRaises(ValueError):parse(workbook(rows))

    def test_misleading_dimensions_do_not_truncate_data(self):
        raw=change_part(workbook(),'xl/worksheets/sheet1.xml',lambda content:content.replace(b'ref="A1:B3"',b'ref="A1:A1"'))
        self.assertEqual(len(parse(raw)['rows']),2)

    def test_actual_row_and_column_limits(self):
        for reference in [b'A10002',b'BM2']:
            raw=change_part(workbook(),'xl/worksheets/sheet1.xml',lambda content:content.replace(b'r="A2"',b'r="'+reference+b'"').replace(b'<row r="2">',b'<row r="10002">') if b'10002' in reference else content.replace(b'r="A2"',b'r="'+reference+b'"'))
            with self.assertRaises(ValueError):parse(raw)

    def test_xml_entities_and_macro_payloads_are_rejected(self):
        raw=change_part(workbook(),'xl/workbook.xml',lambda content:b'<!DOCTYPE x [<!ENTITY x "blocked">]>'+content)
        with self.assertRaisesRegex(ValueError,'XML'):parse(raw)
        target=io.BytesIO(workbook())
        with zipfile.ZipFile(target,'a') as archive:archive.writestr('xl/vbaProject.bin',b'fake')
        with self.assertRaisesRegex(ValueError,'宏'):parse(target.getvalue())

    def test_zip_expansion_budget_and_invalid_format(self):
        raw=change_part(workbook(),'xl/worksheets/sheet1.xml',lambda _:b' '*(8*1024*1024+1))
        with self.assertRaisesRegex(ValueError,'解压'):parse(raw)
        with self.assertRaises(ValueError):parse(b'not excel')

if __name__=='__main__':unittest.main()
